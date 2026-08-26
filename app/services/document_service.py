import os
import shutil
from typing import Optional, List, Tuple
from datetime import datetime
from sqlalchemy.orm import Session
from sqlalchemy import func
from fastapi import UploadFile

from app.core.config import settings
from app.models.postgres_models import (
    Document, DocumentVersion, DocumentTag, Tag, DocumentPermission, User, Role, Department, Category, AuditLog
)
from app.services.chunking_service import chunking_service

class DocumentService:
    def __init__(self, storage_dir: str = settings.STORAGE_DIR):
        self.storage_dir = storage_dir
        os.makedirs(self.storage_dir, exist_ok=True)

    def extract_text_from_file(self, file_path: str) -> str:
        """Extracts plain text content from various file formats (.txt, .md, .pdf, .docx)."""
        ext = os.path.splitext(file_path)[1].lower()
        
        if not os.path.exists(file_path):
            return ""

        try:
            if ext in [".txt", ".md", ".json", ".csv", ".sql", ".py"]:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()

            elif ext == ".pdf":
                try:
                    import pypdf
                    reader = pypdf.PdfReader(file_path)
                    text_parts = [page.extract_text() or "" for page in reader.pages]
                    return "\n".join(text_parts)
                except Exception as e:
                    return f"PDF Extraction Error: {str(e)}"

            elif ext == ".docx":
                try:
                    import docx
                    doc = docx.Document(file_path)
                    return "\n".join([p.text for p in doc.paragraphs if p.text])
                except Exception as e:
                    return f"DOCX Extraction Error: {str(e)}"

            else:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()
        except Exception as e:
            return f"File Reading Error: {str(e)}"

    def check_user_access(self, db: Session, user: User, document: Document) -> Tuple[bool, bool, bool]:
        """
        Determines effective (can_view, can_edit, can_delete) permissions for a user on a document:
        1. Admin: Full View, Edit, Delete access to everything.
        2. Document Owner/Uploader: Full View, Edit, Delete on their own document.
        3. Manager: View and Edit access to documents within their department.
        4. Explicit Grants: Checked from document_permissions table.
        5. Default: Denied.
        """
        if user.role.role_name == "Admin":
            return True, True, True

        if document.uploaded_by == user.user_id:
            return True, True, True

        if user.role.role_name == "Manager" and user.department_id == document.department_id:
            return True, True, False

        perm = db.query(DocumentPermission).filter(
            DocumentPermission.document_id == document.document_id,
            DocumentPermission.user_id == user.user_id
        ).first()

        if perm:
            return bool(perm.can_view), bool(perm.can_edit), bool(perm.can_delete)

        return False, False, False

    def list_accessible_documents(self, db: Session, user: User) -> List[Document]:
        """Returns all documents the specified user has authorization to view."""
        all_docs = db.query(Document).all()
        accessible = []
        for doc in all_docs:
            can_view, _, _ = self.check_user_access(db, user, doc)
            if can_view:
                accessible.append(doc)
        return accessible

    def create_document(
        self,
        db: Session,
        title: str,
        description: str,
        user: User,
        department_id: int,
        category_id: int,
        tags: List[str],
        file_obj: Optional[UploadFile] = None,
        file_content: Optional[str] = None,
        original_filename: Optional[str] = None
    ) -> Document:
        # Determine filename and path
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        if file_obj:
            safe_name = f"{timestamp}_{file_obj.filename.replace(' ', '_')}"
            dest_path = os.path.join(self.storage_dir, safe_name)
            with open(dest_path, "wb") as buffer:
                shutil.copyfileobj(file_obj.file, buffer)
            extracted_text = self.extract_text_from_file(dest_path)
            orig_name = file_obj.filename
        else:
            orig_name = original_filename or f"{title.lower().replace(' ', '_')}.txt"
            safe_name = f"{timestamp}_{orig_name}"
            dest_path = os.path.join(self.storage_dir, safe_name)
            content_to_write = file_content or f"# {title}\n\n{description}"
            with open(dest_path, "w", encoding="utf-8") as f:
                f.write(content_to_write)
            extracted_text = content_to_write

        # 1. Create Document
        doc = Document(
            title=title,
            description=description,
            file_name=orig_name,
            file_path=dest_path,
            uploaded_by=user.user_id,
            department_id=department_id,
            category_id=category_id
        )
        db.add(doc)
        db.flush()

        # 2. Create Initial Version
        v1 = DocumentVersion(
            document_id=doc.document_id,
            version_number=1,
            file_path=dest_path,
            uploaded_by=user.user_id
        )
        db.add(v1)
        db.flush()

        # 3. Associate Tags
        for t_name in tags:
            tag_name_clean = t_name.strip()
            if not tag_name_clean:
                continue
            tag_record = db.query(Tag).filter(Tag.tag_name == tag_name_clean).first()
            if not tag_record:
                tag_record = Tag(tag_name=tag_name_clean)
                db.add(tag_record)
                db.flush()
            db.add(DocumentTag(document_id=doc.document_id, tag_id=tag_record.tag_id))

        # 4. Create Audit Log
        audit = AuditLog(
            user_id=user.user_id,
            action=f"DOCUMENT_UPLOAD: Created document '{title}' (v1)",
            document_id=doc.document_id
        )
        db.add(audit)
        db.commit()

        # 5. Process Text Chunking & Embeddings
        chunking_service.process_and_index_document(
            db=db,
            document_id=doc.document_id,
            text_content=extracted_text,
            version_id=v1.version_id
        )

        db.refresh(doc)
        return doc

    def add_version(
        self,
        db: Session,
        document_id: int,
        user: User,
        file_obj: Optional[UploadFile] = None,
        file_content: Optional[str] = None
    ) -> DocumentVersion:
        doc = db.query(Document).filter(Document.document_id == document_id).first()
        if not doc:
            raise ValueError(f"Document {document_id} not found")

        # Determine next version number
        current_max_v = db.query(func.max(DocumentVersion.version_number)).filter(
            DocumentVersion.document_id == document_id
        ).scalar() or 0
        new_version_num = current_max_v + 1

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        if file_obj:
            safe_name = f"v{new_version_num}_{timestamp}_{file_obj.filename.replace(' ', '_')}"
            dest_path = os.path.join(self.storage_dir, safe_name)
            with open(dest_path, "wb") as buffer:
                shutil.copyfileobj(file_obj.file, buffer)
            extracted_text = self.extract_text_from_file(dest_path)
            doc.file_name = file_obj.filename
        else:
            safe_name = f"v{new_version_num}_{timestamp}_{doc.file_name}"
            dest_path = os.path.join(self.storage_dir, safe_name)
            content_to_write = file_content or f"# {doc.title} (Version {new_version_num})\nUpdated content."
            with open(dest_path, "w", encoding="utf-8") as f:
                f.write(content_to_write)
            extracted_text = content_to_write

        doc.file_path = dest_path
        doc.updated_at = datetime.now()

        new_version = DocumentVersion(
            document_id=document_id,
            version_number=new_version_num,
            file_path=dest_path,
            uploaded_by=user.user_id
        )
        db.add(new_version)
        db.flush()

        audit = AuditLog(
            user_id=user.user_id,
            action=f"DOCUMENT_VERSION_UPLOAD: Added version {new_version_num} for '{doc.title}'",
            document_id=doc.document_id
        )
        db.add(audit)
        db.commit()

        # Re-index chunks with new version text
        chunking_service.process_and_index_document(
            db=db,
            document_id=doc.document_id,
            text_content=extracted_text,
            version_id=new_version.version_id
        )

        return new_version

document_service = DocumentService()
