import io
import uuid
import pytest
from docx import Document as WordDocument
from app.models.postgres_models import User
from app.schemas.search import SearchQuery
from app.schemas.rag import RAGRequest
from app.schemas.nosql import DocumentReviewCreate
from app.schemas.text2sql import Text2SQLRequest
from app.schemas.auth import UserCreate
from app.services.document_service import document_service
from app.services.search_service import search_service
from pydantic import ValidationError

@pytest.mark.parametrize('model,payload', [
    (SearchQuery, {'query': '  '}),
    (SearchQuery, {'query': 'hello', 'top_k': -1}),
    (SearchQuery, {'query': 'hello', 'top_k': 101}),
    (SearchQuery, {'query': 'hello', 'department_id': -1}),
    (RAGRequest, {'question': '  '}),
    (RAGRequest, {'question': 'hello', 'top_k': 0}),
    (DocumentReviewCreate, {'document_id': 1, 'rating': 4, 'review_text': '  '}),
    (Text2SQLRequest, {'natural_query': '  '}),
    (UserCreate, {'name': 'X'*101, 'email': 'valid@example.com', 'password': 'password123', 'role_id': 1, 'department_id': 1}),
])
def test_invalid_inputs_are_rejected(model, payload):
    with pytest.raises(ValidationError):
        model(**payload)

def test_word_table_text_is_extracted(tmp_path):
    word = WordDocument()
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Approved budget'
    table.cell(0, 1).text = 'Finance amount 12000'
    path = tmp_path / 'table.docx'
    word.save(path)
    extracted = document_service.extract_text_from_file(str(path))
    assert 'Approved budget' in extracted
    assert 'Finance amount 12000' in extracted

def test_structured_search_retains_version_for_citations(client, admin_headers, db_session):
    title = 'CitationVersion ' + uuid.uuid4().hex[:8]
    response = client.post('/api/v1/documents', headers=admin_headers, data={
        'title': title, 'department_id': 3, 'category_id': 3, 'content': 'Initial text.'})
    assert response.status_code == 201
    did = response.json()['document_id']
    try:
        updated = client.post(f'/api/v1/documents/{did}/versions', headers=admin_headers,
                              data={'content': 'Current second version citation text.'})
        assert updated.status_code == 200
        results = search_service.structured_search(db_session, title, db_session.get(User, 1))
        found = next(item for item in results if item['document_id'] == did)
        assert found['provenance']['version_number'] == 2
        assert found['provenance']['version_id'] == updated.json()['version_id']
    finally:
        client.delete(f'/api/v1/documents/{did}', headers=admin_headers)

def test_bad_search_inputs_return_422(client, admin_headers):
    response = client.post('/api/v1/search', headers=admin_headers, json={'query': 'policy', 'top_k': -1})
    assert response.status_code == 422


@pytest.mark.parametrize('limit', [-1, 0, 501])
def test_graph_limit_is_validated(client, admin_headers, limit):
    assert client.get('/api/v1/graph', params={'limit': limit}, headers=admin_headers).status_code == 422

def test_graph_entity_lengths_are_validated(client, admin_headers):
    response = client.post('/api/v1/graph/entities', json={'entity_name': 'X'*201, 'entity_type': 'PERSON'}, headers=admin_headers)
    assert response.status_code == 422

def test_admin_review_stats_only_use_existing_documents(client, admin_headers, monkeypatch):
    from app.services.mongo_service import mongo_service
    from app.models.postgres_models import Document
    from app.db.postgres import SessionLocal
    actual = mongo_service.get_telemetry_aggregation
    captured = {}
    def observe(**kwargs):
        captured.update(kwargs)
        return actual(**kwargs)
    monkeypatch.setattr(mongo_service, 'get_telemetry_aggregation', observe)
    response = client.get('/api/v1/nosql/telemetry', headers=admin_headers)
    assert response.status_code == 200
    with SessionLocal() as db:
        existing = {d.document_id for d in db.query(Document).all()}
    assert set(captured['accessible_doc_ids']) == existing
    assert captured['user_id'] is None
